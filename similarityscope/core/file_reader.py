"""
SimilarityScope — File Reader
Extracts text from PDF, DOCX, and TXT files.
"""

import os


def read_pdf(filepath: str) -> tuple[str, str | None]:
    """Extract text from a PDF file using PyMuPDF/fitz."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(filepath)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
        text = "\n".join(text_parts).strip()
        if not text:
            return "", f"PDF file '{os.path.basename(filepath)}' contains no extractable text."
        return text, None
    except Exception as e:
        return "", f"Error reading PDF '{os.path.basename(filepath)}': {str(e)}"


def read_docx(filepath: str) -> tuple[str, str | None]:
    """Extract text from a DOCX file using python-docx.

    Extracts from:
    - Body paragraphs
    - Table cells (merged-cell duplicates skipped via XML element identity)
    - Section headers and footers
    """
    try:
        from docx import Document
        doc = Document(filepath)
        text_parts = []

        # ── Body paragraphs ──
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # ── Tables — deduplicate merged cells by XML element identity ──
        seen_cells = set()
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_id = id(cell._tc)
                    if cell_id not in seen_cells:
                        seen_cells.add(cell_id)
                        if cell.text.strip():
                            text_parts.append(cell.text)

        # ── Section headers and footers ──
        for section in doc.sections:
            for hdr_para in section.header.paragraphs:
                if hdr_para.text.strip():
                    text_parts.append(hdr_para.text)
            for ftr_para in section.footer.paragraphs:
                if ftr_para.text.strip():
                    text_parts.append(ftr_para.text)

        text = "\n".join(text_parts).strip()
        if not text:
            return "", f"DOCX file '{os.path.basename(filepath)}' contains no extractable text."
        return text, None
    except Exception as e:
        return "", f"Error reading DOCX '{os.path.basename(filepath)}': {str(e)}"


def read_txt(filepath: str) -> tuple[str, str | None]:
    """Read text from a plain text file with UTF-8, fallback to latin-1."""
    try:
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                text = f.read().strip()
        except UnicodeDecodeError:
            with open(filepath, 'r', encoding='latin-1') as f:
                text = f.read().strip()
        if not text:
            return "", f"TXT file '{os.path.basename(filepath)}' is empty."
        return text, None
    except Exception as e:
        return "", f"Error reading TXT '{os.path.basename(filepath)}': {str(e)}"


def read_file(filepath: str) -> tuple[str, str | None]:
    """
    Dispatch file reading based on extension.
    Returns (text, error). If error is not None, text will be empty.
    """
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.pdf':
        return read_pdf(filepath)
    elif ext == '.docx':
        return read_docx(filepath)
    elif ext == '.txt':
        return read_txt(filepath)
    else:
        return "", f"Unsupported file type: '{ext}'"
